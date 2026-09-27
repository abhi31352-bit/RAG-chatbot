"""Tests for file parsing and text chunking."""

import io

import pytest

from app.utils.file_handler import (
    ALLOWED_EXTENSIONS,
    DocumentParseError,
    FileHandler,
    UnsupportedFileTypeError,
    normalize_extension,
)
from app.utils.text_processor import (
    InvalidChunkConfigError,
    TextProcessor,
)


def make_docx_bytes(paragraphs, table_rows=None) -> bytes:
    """Build a minimal real .docx in memory."""
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        width = max(len(row) for row in table_rows)
        table = document.add_table(rows=len(table_rows), cols=width)
        for r_index, values in enumerate(table_rows):
            for c_index, value in enumerate(values):
                table.cell(r_index, c_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- Extension handling ----------------------------------------------------


@pytest.mark.parametrize("raw,expected", [
    ("pdf", "pdf"), (".pdf", "pdf"), ("PDF", "pdf"), (" .Docx ", "docx"), ("md", "md"),
])
def test_normalize_extension(raw, expected):
    assert normalize_extension(raw) == expected


def test_normalize_extension_rejects_empty():
    with pytest.raises(UnsupportedFileTypeError):
        normalize_extension("")


def test_allowed_extensions_have_no_dots():
    assert all(not e.startswith(".") for e in ALLOWED_EXTENSIONS)


# -- Text / markdown -------------------------------------------------------


def test_parse_plain_text():
    handler = FileHandler()
    assert handler.parse(b"Hello course world", "txt") == "Hello course world"


def test_parse_markdown():
    handler = FileHandler()
    assert handler.parse(b"# Title\n\nBody", "md").startswith("# Title")


def test_parse_txt_accepts_bare_extension():
    handler = FileHandler()
    assert handler.parse(b"body", ".txt") == "body"


def test_parse_text_handles_utf8_bom():
    handler = FileHandler()
    text = handler.parse(b"\xef\xbb\xbfHello", "txt")
    assert text == "Hello"


def test_parse_text_falls_back_on_latin1():
    handler = FileHandler()
    # 0xE9 is 'e' with an acute accent in latin-1, invalid as UTF-8.
    assert "caf" in handler.parse(b"caf\xe9", "txt")


def test_unsupported_type_raises():
    with pytest.raises(UnsupportedFileTypeError):
        FileHandler().parse(b"data", "exe")


# -- PDF -------------------------------------------------------------------


def test_parse_pdf_and_pages():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    contents = io.BytesIO()
    writer.write(contents)

    handler = FileHandler()
    text, page_boundaries = handler.parse_with_pages(contents.getvalue(), "pdf")
    assert isinstance(text, str)
    assert isinstance(page_boundaries, list)


def test_non_pdf_has_no_page_boundaries():
    _, page_boundaries = FileHandler().parse_with_pages(b"just text", "txt")
    assert page_boundaries == []


def test_parse_corrupt_pdf_raises_parse_error():
    with pytest.raises(DocumentParseError):
        FileHandler().parse(b"this is not a pdf at all", "pdf")


def test_non_pdf_raises_parse_error():
    with pytest.raises(DocumentParseError):
        FileHandler().parse(b"MZ\x00\x00 binary junk", "pdf")


# -- DOCX ------------------------------------------------------------------


def test_parse_docx_paragraphs():
    payload = make_docx_bytes(["Lecture one", "Lecture two"])
    text = FileHandler().parse(payload, "docx")
    assert "Lecture one" in text
    assert "Lecture two" in text


def test_parse_docx_includes_table_cells():
    payload = make_docx_bytes(
        ["Notes"], table_rows=[["Topic", "Week 1"], ["Grading", "40%"]]
    )
    text = FileHandler().parse(payload, "docx")
    assert "Topic | Week 1" in text
    assert "Grading | 40%" in text


def test_parse_non_docx_raises_parse_error():
    with pytest.raises(DocumentParseError):
        FileHandler().parse(b"plain text pretending to be docx", "docx")


# -- Chunking --------------------------------------------------------------


def test_chunk_blank_text_returns_empty():
    assert TextProcessor().chunk_text("   \n  ") == []


def test_chunk_respects_chunk_size():
    processor = TextProcessor()
    text = " ".join(f"token{i}" for i in range(1200))
    chunks = processor.chunk_text(text, chunk_size=100, overlap=10)
    assert len(chunks) > 1
    for chunk in chunks:
        assert processor.count_tokens(chunk) <= 100


def test_chunk_overlap_produces_shared_tokens():
    processor = TextProcessor()
    words = [f"w{i}" for i in range(200)]
    chunks = processor.chunk_text(" ".join(words), chunk_size=50, overlap=25)
    assert len(chunks) >= 3
    first = set(chunks[0].split())
    second = set(chunks[1].split())
    assert first & second, "consecutive chunks should share overlap tokens"


def test_chunk_short_text_is_single_chunk():
    chunks = TextProcessor().chunk_text("just a little text", chunk_size=500, overlap=50)
    assert chunks == ["just a little text"]


def test_chunk_terminates_when_overlap_equals_size():
    # The naive `start = end - overlap` loop would spin forever here.
    with pytest.raises(InvalidChunkConfigError):
        TextProcessor().chunk_text("some text here", chunk_size=50, overlap=50)


def test_chunk_rejects_overlap_larger_than_size():
    with pytest.raises(InvalidChunkConfigError):
        TextProcessor().chunk_text("some text here", chunk_size=50, overlap=80)


def test_chunk_rejects_non_positive_size():
    with pytest.raises(InvalidChunkConfigError):
        TextProcessor().chunk_text("text", chunk_size=0)


def test_chunk_covers_entire_document():
    processor = TextProcessor()
    words = [f"unique{i}" for i in range(300)]
    chunks = processor.chunk_text(" ".join(words), chunk_size=80, overlap=20)
    joined = " ".join(chunks)
    assert "unique0" in joined
    assert "unique299" in joined


def test_count_tokens_empty_is_zero():
    assert TextProcessor().count_tokens("") == 0


# -- Chunk spans (used to attribute chunks to pages) ----------------------


def test_spans_match_chunk_text_output():
    processor = TextProcessor()
    text = " ".join(f"word{i}" for i in range(400))
    plain = processor.chunk_text(text, chunk_size=60, overlap=10)
    spanned = processor.chunk_text_with_spans(text, chunk_size=60, overlap=10)
    assert [c for c, _, _ in spanned] == plain


def test_spans_are_ordered_and_within_bounds():
    processor = TextProcessor()
    text = " ".join(f"token{i}" for i in range(500))
    spans = processor.chunk_text_with_spans(text, chunk_size=80, overlap=20)

    assert spans
    previous_start = -1
    for chunk, start, end in spans:
        assert 0 <= start <= end <= len(text)
        assert start > previous_start, "chunk starts should advance"
        previous_start = start
        # The span should really point at this chunk's text.
        assert text[start:end].strip()[:20] == chunk[:20]


def test_spans_first_chunk_starts_at_zero():
    processor = TextProcessor()
    text = "alpha beta gamma delta epsilon"
    _, start, _ = processor.chunk_text_with_spans(text, chunk_size=5, overlap=1)[0]
    assert start == 0


def test_spans_cover_later_characters():
    """Regression: page attribution needs spans that advance through the text."""
    processor = TextProcessor()
    text = " ".join(f"sentence number {i} about machine learning" for i in range(200))
    spans = processor.chunk_text_with_spans(text, chunk_size=100, overlap=10)
    assert len(spans) > 3
    assert spans[-1][1] > 0, "later chunks must start past the beginning of the text"
